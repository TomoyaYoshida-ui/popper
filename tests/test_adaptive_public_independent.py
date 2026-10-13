import io

import numpy as np
import openpyxl
import pytest

from popper.core import ProtocolError
from popper.research.adaptive_public_independent import (
    materialize_task,
    parse_concrete_table,
    parse_energy_table,
    parse_forest_text,
    parse_qsar_text,
    parse_servo_text,
    parse_yacht_text,
)


def test_concrete_skips_header_and_keeps_eight_inputs():
    rows = [("Cement", "Slag", "Strength"), (540, 0, 1, 162, 2.5, 1040, 676, 28, 80)]
    features, target, names = parse_concrete_table(rows)
    assert features.shape == (1, 8)
    assert target.tolist() == [80]
    assert names[-1] == "age"


def test_yacht_reads_whitespace_rows():
    features, target, names = parse_yacht_text("-2.3 0.568 4.78 3.99 3.17 0.125 0.11\n")
    assert features.shape == (1, 6)
    assert target.tolist() == [0.11]
    assert "froude" in names


def test_energy_uses_heating_load_not_cooling_load():
    features, target, names = parse_energy_table([(0.98, 514, 294, 110, 7, 2, 0, 0, 15.55, 21.33, None, None)])
    assert list(features[0])[:1] == [0.98]
    assert target.tolist() == [15.55]
    assert "heating" not in names[-1] or names[-1] == "glazing_distribution"
    assert features.shape[1] == 8


def test_qsar_accepts_semicolons():
    features, target, names = parse_qsar_text("1;2;3;4;5;6;7;8;0.5\n")
    assert features.shape == (1, 8)
    assert target.tolist() == [0.5]
    assert names[0] == "tpsa"


def test_servo_one_hot_encodes_linkages():
    features, target, names = parse_servo_text("E,A,5,4,0.28125095\n")
    assert features.shape == (1, 12)
    assert features[0, 4] == 1 and features[0, 5] == 1
    assert target.tolist() == [0.28125095]
    assert names[0] == "motor_A" and names[-1] == "vgain"


def test_forest_one_hot_encodes_month_and_day():
    header = "X,Y,month,day,FFMC,DMC,DC,ISI,temp,RH,wind,rain,area"
    text = header + "\n7,5,mar,fri,86.2,26.2,94.3,5.1,8.2,51,6.7,0,0\n"
    features, target, names = parse_forest_text(text)
    assert features.shape == (1, 2 + 12 + 7 + 8)
    assert features[0, names.index("month_mar")] == 1
    assert features[0, names.index("day_fri")] == 1
    assert target.tolist() == [0]


def test_unknown_servo_level_is_rejected():
    with pytest.raises(ProtocolError):
        parse_servo_text("Z,A,5,4,0.2\n")


def test_materialize_marks_public_independence_unverified(tmp_path):
    features = np.array([[1.0, 2.0], [3.0, 4.0], [5.0, 6.0], [7.0, 8.0], [9.0, 0.0]])
    target = np.array([1.0, 2.0, 3.0, 4.0, 5.0])
    prepared = materialize_task(
        tmp_path / "task", "U01", features, target, ["a", "b"],
        {"kind": "public", "dataset_id": 165, "target_column": "compressive_strength"})
    assert prepared["independence_verified"] is False
    assert prepared["sampling"] == {"design": "feature_row_hash"}
    assert prepared["source"]["dataset_id"] == 165
    assert set(prepared["data_hashes"]) == {"train", "dev", "confirmation"}


def test_energy_xlsx_round_trip_table():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["X1", "X2", "X3", "X4", "X5", "X6", "X7", "X8", "Y1", "Y2"])
    sheet.append([0.98, 514, 294, 110, 7, 2, 0, 0, 15.55, 21.33])
    payload = io.BytesIO()
    book.save(payload)
    restored = openpyxl.load_workbook(io.BytesIO(payload.getvalue()), data_only=True)
    rows = list(restored.active.iter_rows(values_only=True))
    features, target, _names = parse_energy_table(rows)
    assert target.tolist() == [15.55]
    assert features.shape == (1, 8)
